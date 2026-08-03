from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl


_REPO_ROOT = Path(__file__).resolve().parents[4]

DEFAULT_TXT_DIR = _REPO_ROOT / "archive/legacy_research/assets/pager_db/txt"
DEFAULT_XLSX_PATH = _REPO_ROOT / "archive/legacy_research/assets/pager_db/pager_db_merged_最终标注版.xlsx"
DEFAULT_MANIFEST_PATH = Path("outputs/local_retrieval/corpus_manifest.jsonl")

MANIFEST_METADATA_ALIASES: dict[str, list[str]] = {
    "pdf_path": ["文件路径", "source_file", "source_file_来源 2"],
    "manual_score": ["人工打分_来源 2", "人工评分", "人工打分", "评分"],
    "relevance_to_research_score": [
        "relevance_to_research_score_来源 2",
        "relevance_to_research_score",
        "relevance_score",
    ],
    "authors": ["authors_来源 2", "authors"],
    "publication_date": ["publication_date_来源 2", "publication_date"],
    "source": ["source_来源 2", "source"],
    "document_type": ["document_type_来源 2", "document_type"],
}


def normalize_doc_record(txt_path: str) -> dict[str, str]:
    path = Path(txt_path)
    title = path.stem
    return {
        "doc_id": title.lower().replace(" ", "_"),
        "title": title,
        "txt_path": str(path),
    }


def _first_non_empty(row: dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        value = row.get(alias)
        if value not in (None, ""):
            return value
    return None


def load_metadata_rows(xlsx_path: str | Path = DEFAULT_XLSX_PATH) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))

    rows: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append({headers[index]: row[index] for index in range(len(headers))})
    return rows


def build_metadata_lookup(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    lookup: dict[str, dict[str, Any]] = {}
    for row in rows:
        filename = row.get("文件名")
        if filename in (None, ""):
            continue
        key = Path(str(filename)).stem.lower().replace(" ", "_")
        lookup[key] = row
    return lookup


def merge_metadata(record: dict[str, str], metadata_row: dict[str, Any] | None) -> dict[str, Any]:
    merged: dict[str, Any] = dict(record)
    merged["language"] = "zh" if any("\u4e00" <= char <= "\u9fff" for char in record["title"]) else "en"
    merged["source_type"] = "txt_extract"
    merged["pdf_path"] = None
    for field_name in MANIFEST_METADATA_ALIASES:
        merged[field_name] = None

    if not metadata_row:
        return merged

    for field_name, aliases in MANIFEST_METADATA_ALIASES.items():
        merged[field_name] = _first_non_empty(metadata_row, aliases)
    if merged.get("pdf_path") in (None, ""):
        merged["pdf_path"] = None
    return merged


def build_corpus_manifest(
    txt_dir: str | Path = DEFAULT_TXT_DIR,
    xlsx_path: str | Path = DEFAULT_XLSX_PATH,
) -> list[dict[str, Any]]:
    metadata_lookup = build_metadata_lookup(load_metadata_rows(xlsx_path))
    manifest: list[dict[str, Any]] = []

    for txt_file in sorted(Path(txt_dir).glob("*.txt")):
        record = normalize_doc_record(str(txt_file))
        metadata_row = metadata_lookup.get(record["doc_id"])
        manifest.append(merge_metadata(record, metadata_row))

    return manifest


def write_corpus_manifest(
    output_path: str | Path = DEFAULT_MANIFEST_PATH,
    txt_dir: str | Path = DEFAULT_TXT_DIR,
    xlsx_path: str | Path = DEFAULT_XLSX_PATH,
) -> list[dict[str, Any]]:
    manifest = build_corpus_manifest(txt_dir=txt_dir, xlsx_path=xlsx_path)
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with output_file.open("w", encoding="utf-8") as handle:
        for row in manifest:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    return manifest
